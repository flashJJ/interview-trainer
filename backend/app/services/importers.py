"""多格式资料导入解析：PDF / DOCX / 网页链接。

- PDF：pypdf 提取文本层（扫描件无文本层时明确报错，不做 OCR）
- DOCX：python-docx 提取段落文本
- 网页：httpx 抓取 + readability-lxml 提取正文（去广告/导航/页脚）
"""

import io
import re

import httpx
from docx import Document
from pypdf import PdfReader
from readability import Document as ReadabilityDocument

MAX_URL_BYTES = 10 * 1024 * 1024  # 网页内容上限 10MB
URL_TIMEOUT = 30  # 秒

# 明显的非正文噪声（readability 偶尔漏掉的导航残留）
_NOISE_LINE_RE = re.compile(r"^(登录|注册|首页|首页\s*\|\s*.*|版权所有|copyright|all rights reserved)", re.IGNORECASE)


def extract_pdf_text(raw: bytes) -> str:
    """提取 PDF 文本层内容。"""
    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"PDF 解析失败（文件可能损坏或非 PDF 格式）：{exc}") from exc
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # 尝试空密码
        except Exception as exc:  # noqa: BLE001
            raise ValueError("PDF 已加密，无法提取文本") from exc
    parts = []
    for page in reader.pages[:200]:  # 上限 200 页，防止超大文件
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - 单页失败不阻塞整体
            continue
    text = "\n".join(parts)
    if len(text.strip()) < 30:
        raise ValueError(
            "未能从 PDF 提取到文本（可能是扫描件/图片型 PDF，暂不支持 OCR）"
        )
    return text


def extract_docx_text(raw: bytes) -> str:
    """提取 DOCX 段落文本。"""
    try:
        document = Document(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"DOCX 解析失败（文件可能损坏）：{exc}") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # 表格内容也提取
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if len(text.strip()) < 30:
        raise ValueError("DOCX 内容为空或过少")
    return text


def extract_html_text(html: str) -> str:
    """用 readability 从 HTML 提取正文。"""
    if not html or not html.strip():
        raise ValueError("HTML 内容为空")
    try:
        document = ReadabilityDocument(html)
        title = document.short_title()
        content_html = document.summary(html_partial=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"网页正文提取失败：{exc}") from exc
    # 去标签得到纯文本
    text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>", " ", content_html, flags=re.I)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</(p|div|li|h[1-6]|tr)>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = (
        text.replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line and not _NOISE_LINE_RE.match(line)]
    body = "\n".join(lines)
    if len(body) < 30:
        raise ValueError("网页正文提取结果过少（可能需要登录或为动态渲染页面）")
    return f"# {title}\n\n{body}" if title else body


def fetch_url_text(url: str) -> str:
    """抓取网页并提取正文。"""
    if not re.match(r"^https?://", url.strip(), re.IGNORECASE):
        raise ValueError("URL 必须以 http:// 或 https:// 开头")
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,*/*;q=0.8",
    }
    try:
        # trust_env=False：直连目标地址，避免代理拦截内网/本地地址
        with httpx.Client(timeout=URL_TIMEOUT, follow_redirects=True, max_redirects=5, trust_env=False) as client:
            resp = client.get(url.strip(), headers=headers)
            resp.raise_for_status()
            content_length = int(resp.headers.get("content-length") or 0)
            if content_length > MAX_URL_BYTES:
                raise ValueError("网页内容超过 10MB 上限")
            html = resp.text
    except httpx.HTTPStatusError as exc:
        raise ValueError(f"网页返回错误状态 {exc.response.status_code}") from exc
    except httpx.TimeoutException as exc:
        raise ValueError(f"抓取超时（{URL_TIMEOUT}s）") from exc
    except httpx.TooManyRedirects as exc:
        raise ValueError("重定向次数过多") from exc
    except httpx.RequestError as exc:
        raise ValueError(f"网络请求失败：{exc}") from exc
    return extract_html_text(html)
